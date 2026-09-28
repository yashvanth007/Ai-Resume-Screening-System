import re
from pathlib import Path

import fitz
from docx import Document

from app.ml.skills import ALL_SKILLS, ALIASES, SKILL_CATEGORY, SKILL_DISPLAY

SECTION_HEADERS = {
    "summary": ("summary", "profile", "objective", "about me"),
    "skills": ("skills", "technical skills", "core competencies", "competencies"),
    "experience": ("experience", "work experience", "employment", "professional experience"),
    "education": ("education", "academic background", "qualifications"),
    "projects": ("projects", "personal projects", "academic projects"),
    "certifications": ("certifications", "certificates", "licenses"),
}


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document)
    elif suffix == ".docx":
        import io
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raise ValueError("Only PDF and DOCX resumes are supported.")
    cleaned = re.sub(r"[ \t]+", " ", text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if len(cleaned) < 30:
        raise ValueError("This file contains too little readable text. Scanned PDFs need OCR before upload.")
    return cleaned


def extract_sections(text: str) -> dict[str, str]:
    lines = text.splitlines()
    markers: list[tuple[int, str]] = []
    for index, line in enumerate(lines):
        normalized = re.sub(r"[^a-z ]", "", line.casefold()).strip()
        normalized = re.sub(r"\s+", " ", normalized)
        for section, aliases in SECTION_HEADERS.items():
            if normalized in aliases:
                markers.append((index, section))
                break
    sections: dict[str, str] = {}
    for position, (line_index, section) in enumerate(markers):
        end = markers[position + 1][0] if position + 1 < len(markers) else len(lines)
        sections[section] = "\n".join(lines[line_index + 1:end]).strip()
    return sections


def extract_skills(text: str) -> list[dict[str, str]]:
    found: dict[str, dict[str, str]] = {}
    folded = text.casefold()
    for raw_skill in ALL_SKILLS:
        escaped = re.escape(raw_skill)
        if re.search(rf"(?<![\w]){escaped}(?![\w])", folded, flags=re.IGNORECASE):
            display_name = SKILL_DISPLAY[raw_skill]
            canonical = ALIASES.get(display_name, display_name)
            found[canonical.casefold()] = {"name": canonical, "category": SKILL_CATEGORY[raw_skill]}
    return sorted(found.values(), key=lambda item: item["name"].casefold())


def _first(pattern: str, text: str, flags: int = re.IGNORECASE) -> str | None:
    match = re.search(pattern, text, flags)
    return match.group(1).strip(" .,;|") if match else None


def parse_resume(text: str) -> dict:
    sections = extract_sections(text)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email = _first(r"\b([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})\b", text)
    phone = _first(r"(?<!\w)(\+?\d[\d\s().-]{7,}\d)(?!\w)", text)
    links = re.findall(r"https?://[^\s<>]+", text, re.IGNORECASE)
    name = next((line for line in lines[:8] if len(line) <= 80 and not re.search(r"@|https?://|\d{3}", line) and not any(word in line.casefold() for word in ("resume", "curriculum vitae", "linkedin", "github"))), None)
    years_match = re.search(r"(\d+(?:\.\d+)?)\s*\+?\s+years?", text, re.IGNORECASE)
    experience_years = float(years_match.group(1)) if years_match else 0.0
    skill_items = extract_skills(text)
    education_lines = [line for line in sections.get("education", "").splitlines() if line.strip()]
    project_lines = [line.strip(" •-*\t") for line in sections.get("projects", "").splitlines() if line.strip(" •-*\t")]
    certification_lines = [line.strip(" •-*\t") for line in sections.get("certifications", "").splitlines() if line.strip(" •-*\t")]
    gpa = _first(r"\b(?:CGPA|GPA)\s*[:=-]?\s*(\d+(?:\.\d+)?(?:\s*/\s*\d+)?)", text)
    year_match = re.search(r"\b(19\d{2}|20[0-3]\d)\b", sections.get("education", ""))
    degree = _first(r"\b(Bachelor(?:'s)?|Master(?:'s)?|Ph\.?D\.?|B\.?Tech|M\.?Tech|B\.?Sc|M\.?Sc|BCA|MCA|MBA|Associate(?:'s)?)\b[^\n,;]{0,70}", sections.get("education", ""))
    location = next((line for line in lines[:10] if re.search(r"\b(?:India|United States|USA|UK|Canada|Australia|Germany|Remote)\b", line, re.IGNORECASE) and not re.search(r"@|http", line)), None)
    return {
        "name": name or "Not detected", "email": email, "phone": phone, "location": location,
        "linkedin": next((url for url in links if "linkedin.com" in url.casefold()), None),
        "github": next((url for url in links if "github.com" in url.casefold()), None),
        "portfolio": next((url for url in links if "linkedin.com" not in url.casefold() and "github.com" not in url.casefold()), None),
        "experience_years": experience_years,
        "skills": skill_items,
        "education": [{"degree": degree, "institution": education_lines[0] if education_lines else None, "graduation_year": int(year_match.group(1)) if year_match else None, "gpa": gpa}] if education_lines or degree else [],
        "projects": [{"name": line[:120], "description": line, "technologies": [item["name"] for item in skill_items if re.search(rf"(?<![\w]){re.escape(item['name'])}(?![\w])", line, re.IGNORECASE)]} for line in project_lines[:12]],
        "certifications": [{"name": line, "issuer": None} for line in certification_lines[:12]],
        "sections": sections,
    }
