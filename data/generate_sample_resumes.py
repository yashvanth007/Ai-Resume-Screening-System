from pathlib import Path

from docx import Document

OUTPUT = Path(__file__).parent / "sample_resumes"
PROFILES = [
    ("Avery Chen", "avery.chen@example.test", "Portland, OR", "4 years experience", ["Python", "FastAPI", "SQL", "REST API", "Git", "Docker", "AWS", "PostgreSQL"], "B.S. Computer Science", "Cascadia Technical University", "Service observability platform", "Built a Python FastAPI platform with PostgreSQL, Docker, and AWS deployment pipelines."),
    ("Jordan Rivera", "jordan.rivera@example.test", "Chicago, IL", "3 years experience", ["Python", "Machine Learning", "scikit-learn", "NLP", "Pandas", "PyTorch", "Git"], "M.S. Data Science", "Lakeview Institute", "Support ticket intent classifier", "Trained and evaluated an NLP classifier with scikit-learn and PyTorch on support text."),
    ("Samira Okafor", "samira.okafor@example.test", "Atlanta, GA", "5 years experience", ["Python", "SQL", "Pandas", "Machine Learning", "Power BI", "AWS", "Git"], "B.S. Applied Mathematics", "Piedmont State College", "Product adoption analysis", "Analyzed product funnels with SQL and pandas and delivered Power BI reporting."),
    ("Taylor Brooks", "taylor.brooks@example.test", "Denver, CO", "2 years experience", ["Python", "FastAPI", "SQL", "REST API", "Git", "Machine Learning", "scikit-learn"], "B.S. Software Engineering", "Front Range University", "Inventory demand forecasting", "Created Python forecasting models and served them through FastAPI and SQL."),
    ("Riley Morgan", "riley.morgan@example.test", "Boston, MA", "1 year experience", ["Python", "Pandas", "SQL", "Machine Learning", "Git", "TensorFlow"], "B.S. Statistics", "Commonwealth College", "Transit delay prediction", "Compared TensorFlow and baseline models with Python, pandas, and SQL."),
]


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for index, profile in enumerate(PROFILES, 1):
        name, email, location, experience, skills, degree, school, project, description = profile
        doc = Document()
        doc.add_heading(name, 0)
        doc.add_paragraph(f"{email} | {location}")
        doc.add_paragraph(experience)
        for section, lines in [("Summary", [f"Early-career professional with experience in {', '.join(skills[:4])}."]), ("Skills", [", ".join(skills)]), ("Experience", [f"Software Engineer | Sample Organization | {description}"]), ("Education", [degree, school]), ("Projects", [project, description]), ("Certifications", ["Cloud Foundations Certificate | Sample Learning Institute"])]:
            doc.add_heading(section, level=1)
            for line in lines:
                doc.add_paragraph(line)
        doc.save(OUTPUT / f"sample_resume_{index:02}.docx")
    print(f"Generated {len(PROFILES)} synthetic DOCX resumes in {OUTPUT}")


if __name__ == "__main__":
    main()
